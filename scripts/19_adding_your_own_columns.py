# A spreadsheet is not limited to what the database gives you.
# You can add any columns you like: today's date, a calculation, a blank
# column for someone to fill in, a flag, or even a random sample number.
import random
from datetime import date

import mysql.connector
from openpyxl import Workbook
from openpyxl.styles import Font

start = "2026-06-01"
end = "2026-06-30"

conn = mysql.connector.connect(
    host="127.0.0.1", user="aquacertify", password="aquacertify",
    database="aquacertify", charset="utf8mb4", collation="utf8mb4_general_ci",
)
cursor = conn.cursor(dictionary=True)
cursor.execute("""
    SELECT wells.name AS well,
           well_meters.fixed_zeros,
           COUNT(*) AS readings,
           ROUND(AVG(chlorine_residual), 2) AS avg_chlorine,
           MAX(meter_reading) - MIN(meter_reading) AS dial_difference
    FROM well_readings
    JOIN wells       ON wells.id = well_readings.well_id
    JOIN well_meters ON well_meters.id = well_readings.meter_id
    WHERE wells.organization_id = %s AND date_of_reading BETWEEN %s AND %s
      AND well_meters.active = 1
    GROUP BY wells.name, well_meters.fixed_zeros
    ORDER BY wells.name
""", (1, start, end))
rows = cursor.fetchall()
conn.close()

workbook = Workbook()
sheet = workbook.active
sheet.title = "June wells"

# Columns from the database, PLUS columns we invent ourselves:
sheet.append([
    "Well", "Readings", "Avg chlorine",      # straight from the database
    "Gallons",                               # calculated in Python
    "Chlorine OK?",                          # a yes/no flag we decide
    "Report date",                           # today's date
    "Checked by",                            # left blank for a person to fill in
    "Sample #",                              # a random number, e.g. for a spot-check lottery
])

for row in rows:
    gallons = row["dial_difference"] * 10 ** row["fixed_zeros"]
    if row["avg_chlorine"] is not None and row["avg_chlorine"] >= 0.5:
        chlorine_ok = "yes"
    else:
        chlorine_ok = "CHECK"
    sheet.append([
        row["well"], row["readings"], row["avg_chlorine"],
        gallons,
        chlorine_ok,
        date.today(),
        "",
        random.randint(1000, 9999),
    ])

# A total row at the bottom, using a real Excel formula so it updates if someone edits the sheet.
last = sheet.max_row
sheet.append(["TOTAL", "", "", f"=SUM(D2:D{last})"])
sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)

for cell in sheet[1]:
    cell.font = Font(bold=True)
for r in range(2, sheet.max_row + 1):
    sheet.cell(row=r, column=4).number_format = "#,##0"          # 1,234,567
    sheet.cell(row=r, column=6).number_format = "yyyy-mm-dd"
sheet.column_dimensions["A"].width = 24
sheet.column_dimensions["D"].width = 16

workbook.save("output/june_wells_with_extras.xlsx")
print("Saved output/june_wells_with_extras.xlsx")
