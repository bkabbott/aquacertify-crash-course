# A full monthly report: three tabs and a chart, all from the database.
from datetime import date

import mysql.connector
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill

start = "2026-06-01"
end = "2026-06-30"

conn = mysql.connector.connect(
    host="127.0.0.1", user="aquacertify", password="aquacertify",
    database="aquacertify", charset="utf8mb4", collation="utf8mb4_general_ci",
)
cursor = conn.cursor(dictionary=True)

# --- question 1: wells -------------------------------------------------------
cursor.execute("""
    SELECT wells.name AS well, systems.name AS system_name,
           COUNT(*) AS readings,
           ROUND(AVG(chlorine_residual), 2) AS avg_chlorine,
           CAST((MAX(meter_reading) - MIN(meter_reading)) * POW(10, well_meters.fixed_zeros) AS UNSIGNED) AS gallons
    FROM well_readings
    JOIN wells       ON wells.id = well_readings.well_id
    JOIN systems     ON systems.id = wells.system_id
    JOIN well_meters ON well_meters.id = well_readings.meter_id
    WHERE wells.organization_id = %s AND date_of_reading BETWEEN %s AND %s
      AND well_meters.active = 1
    GROUP BY wells.name, systems.name, well_meters.fixed_zeros
    ORDER BY gallons DESC
""", (1, start, end))
wells = cursor.fetchall()

# --- question 2: lift stations ----------------------------------------------
cursor.execute("""
    SELECT lift_stations.name AS station,
           ROUND(MAX(pump1) - MIN(pump1), 1) AS pump1_hours,
           ROUND(MAX(pump2) - MIN(pump2), 1) AS pump2_hours,
           ROUND(SUM(rainfall), 2) AS rain_inches
    FROM station_readings
    JOIN lift_stations ON lift_stations.id = station_readings.lift_station_id
    WHERE lift_stations.organization_id = %s AND date_of_reading BETWEEN %s AND %s
    GROUP BY lift_stations.name
    ORDER BY lift_stations.name
""", (1, start, end))
stations = cursor.fetchall()

# --- question 3: lead addresses still to check --------------------------------
cursor.execute("""
    SELECT routes.name AS route, sli_ga_addresses.street_address_1 AS address,
           sli_ga_addresses.city, sli_ga_addresses.zip
    FROM sli_ga_addresses
    JOIN routes ON routes.id = sli_ga_addresses.route_id
    JOIN sli_ga_service_lines ON sli_ga_service_lines.sli_ga_address_id = sli_ga_addresses.id
    JOIN sli_ga_overall_service_line_classification_responses AS answers
      ON answers.id = sli_ga_service_lines.overall_service_line_classification_id
    WHERE sli_ga_addresses.organization_id = %s AND answers.response = 'Lead Status Unknown'
    ORDER BY routes.name, sli_ga_addresses.street_address_1
""", (1,))
to_check = cursor.fetchall()
conn.close()


def fill_sheet(sheet, headers, rows, columns):
    """Write headers + rows onto a sheet and make the header row bold on blue.
    `columns` is the list of dictionary keys, in the order the headers are in."""
    sheet.append(headers)
    for row in rows:
        sheet.append([row[c] for c in columns])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    for i, header in enumerate(headers):
        sheet.column_dimensions[chr(ord("A") + i)].width = max(12, len(header) + 2)


workbook = Workbook()

# Tab 1: wells + total + chart
sheet = workbook.active
sheet.title = "Wells"
fill_sheet(sheet, ["Well", "System", "Readings", "Avg chlorine", "Gallons"], wells,
           ["well", "system_name", "readings", "avg_chlorine", "gallons"])
sheet.column_dimensions["A"].width = 24
sheet.column_dimensions["B"].width = 30
sheet.append([])
sheet.append(["TOTAL", "", "", "", sum(w["gallons"] for w in wells)])
for r in range(2, sheet.max_row + 1):
    sheet.cell(row=r, column=5).number_format = "#,##0"

chart = BarChart()
chart.title = f"Gallons pumped, {start} to {end}"
top = min(len(wells), 10)
chart.add_data(Reference(sheet, min_col=5, min_row=1, max_row=top + 1), titles_from_data=True)
chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=top + 1))
chart.width, chart.height = 24, 12
sheet.add_chart(chart, "H2")

# Tab 2 and 3
fill_sheet(workbook.create_sheet("Lift stations"), ["Station", "Pump 1 hours", "Pump 2 hours", "Rain (in)"],
           stations, ["station", "pump1_hours", "pump2_hours", "rain_inches"])
fill_sheet(workbook.create_sheet("Lead - to check"), ["Route", "Address", "City", "Zip"],
           to_check, ["route", "address", "city", "zip"])

filename = f"output/monthly_report_{start[:7]}.xlsx"
workbook.save(filename)
print(f"Saved {filename}: {len(wells)} wells, {len(stations)} stations, {len(to_check)} addresses to check")
