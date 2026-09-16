# The same list, but with bold headers, sensible column widths and a frozen header row.
import mysql.connector
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

conn = mysql.connector.connect(
    host="127.0.0.1", user="aquacertify", password="aquacertify",
    database="aquacertify", charset="utf8mb4", collation="utf8mb4_general_ci",
)
cursor = conn.cursor(dictionary=True)
cursor.execute("""
    SELECT wells.name AS well, routes.name AS route, wells.latitude, wells.longitude
    FROM wells
    JOIN routes ON routes.id = wells.route_id
    WHERE wells.organization_id = %s
    ORDER BY routes.name, wells.name
""", (1,))
rows = cursor.fetchall()
conn.close()

workbook = Workbook()
sheet = workbook.active
sheet.title = "Wells"

sheet.append(["Well", "Route", "Latitude", "Longitude"])
for row in rows:
    sheet.append([row["well"], row["route"], row["latitude"], row["longitude"]])

# --- styling -------------------------------------------------------------
for cell in sheet[1]:                                   # every cell in row 1
    cell.font = Font(bold=True, color="FFFFFF")         # white bold text
    cell.fill = PatternFill("solid", fgColor="1F4E78")  # dark blue background

sheet.column_dimensions["A"].width = 24
sheet.column_dimensions["B"].width = 14
sheet.column_dimensions["C"].width = 12
sheet.column_dimensions["D"].width = 12

sheet.freeze_panes = "A2"                               # headers stay put when you scroll

workbook.save("output/wells_styled.xlsx")
print("Saved output/wells_styled.xlsx")
